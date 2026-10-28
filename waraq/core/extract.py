"""Arabic-safe text extraction: PDF -> logical Unicode text / Word.

pdfminer.six (MIT) gives every glyph with its position.  We then

1. map presentation-form letters to normal letters (``ﺑ`` -> ``ب``),
2. sort the glyphs of each line by their x position (so the result does
   not depend on the order the PDF producer wrote them),
3. attach harakat to the letter under them,
4. rebuild word spaces from the gaps between glyphs, and
5. undo the bidi reordering (visual -> logical) per line.
"""
from __future__ import annotations

import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable, List, Optional

from .arabic import has_arabic, is_arabic_char, is_mark, normalize_presentation, visual_units_to_logical
from .errors import NeedsPasswordError, WaraqError, WrongPasswordError


@dataclass
class Line:
    text: str
    x0: float
    y0: float
    x1: float
    y1: float
    size: float
    rtl: bool
    units: List[str] = field(default_factory=list, repr=False)


@dataclass
class Block:
    lines: List[Line]
    rtl: bool

    @property
    def text(self) -> str:
        return " ".join(l.text for l in self.lines if l.text)


@dataclass
class PageText:
    number: int                      # 1-based
    blocks: List[Block] = field(default_factory=list)
    chars: int = 0
    has_images: bool = False
    ocr: bool = False                # filled by OCR instead of the text layer

    @property
    def needs_ocr(self) -> bool:
        return self.chars < 15 and self.has_images

    def lines(self) -> List[str]:
        return [l.text for b in self.blocks for l in b.lines]


def _char_text(c) -> str:
    t = c.get_text()
    if t.startswith("(cid:"):
        return ""
    return normalize_presentation(t)


def _line_units(chars) -> (List[str], float):
    """Order glyphs of one line left->right, attach marks, add spaces."""
    items = []
    for seq, c in enumerate(chars):
        t = _char_text(c)
        if not t:
            continue
        items.append((c.x0, c.x1, t, c.size, getattr(c, "waraq_seq", seq)))
    if not items:
        return [], 0.0
    items = _drop_unmapped(items)
    bases = [it for it in items if not all(is_mark(ch) for ch in it[2])]
    marks = [it for it in items if all(is_mark(ch) for ch in it[2])]
    # centre, not left edge: some producers (MS Word) overlap a space with a
    # wide final letter, and only the centres keep the drawing order
    bases.sort(key=lambda it: (round((it[0] + it[1]) / 2, 1), it[4]))
    # two glyphs drawn on top of each other (space + letter): their x order is
    # meaningless, keep the order they were drawn in
    for k in range(len(bases) - 1):
        a, b = bases[k], bases[k + 1]
        ov = min(a[1], b[1]) - max(a[0], b[0])
        narrow = max(min(a[1] - a[0], b[1] - b[0]), 0.01)
        if ov > 0.6 * narrow and a[4] > b[4]:
            bases[k], bases[k + 1] = b, a
    texts = [it[2] for it in bases]
    # attach each mark to the base glyph it sits on
    for mx0, mx1, mt, _s, _q in marks:
        cx = (mx0 + mx1) / 2
        best, bestd = None, None
        for k, (bx0, bx1, bt, _bs, _bq) in enumerate(bases):
            if bt.isspace():
                continue
            # a zero-width mark drawn exactly at a glyph boundary belongs to the glyph
            # that starts there (reversed-string producers draw it before its base)
            d = -1.0 if bx0 - 0.01 <= cx < bx1 - 0.01 else min(abs(cx - bx0), abs(cx - bx1))
            if bestd is None or d < bestd:
                best, bestd = k, d
        if best is not None:
            texts[best] += mt
    size = statistics.median([it[3] for it in bases]) if bases else 10.0
    units: List[str] = []
    prev = None
    for k, it in enumerate(bases):
        if prev is not None:
            gap = it[0] - prev[1]
            if (gap > size * 0.15 and not it[2].isspace() and not prev[2].isspace()
                    and (not units or units[-1] != " ")):
                units.append(" ")
        t = texts[k]
        if t.isspace():
            if units and units[-1] == " ":
                prev = it
                continue
            t = " "
        units.append(t)
        prev = it
    while units and units[0] == " ":
        units.pop(0)
    while units and units[-1] == " ":
        units.pop()
    return units, size


def _drop_unmapped(items):
    """Remove glyphs that have no real Unicode mapping.

    Producers such as LibreOffice/Word leave some shaping glyphs (kashida
    used for justification, mark variants, ligature connectors) without a
    usable ToUnicode entry; PDF readers then show a random char such as
    "%", ".", "N", "°" or "¶".  Signs of that: a zero-width non-mark glyph,
    or a run of non-Arabic glyphs glued between two Arabic letters.
    """
    import unicodedata

    def arabic_letter(t):
        return any(is_arabic_char(ch) and unicodedata.category(ch).startswith("L") for ch in t)

    items = [it for it in items
             if all(is_mark(ch) for ch in it[2]) or it[1] - it[0] >= 0.01 or it[2].isspace()]
    order = sorted(range(len(items)), key=lambda i: (items[i][0], items[i][4]))
    seq = [items[i] for i in order]
    drop = set()
    bridge = []
    k = 0
    while k < len(seq):
        x0, x1, t, size, _ = seq[k]
        suspicious = (not t.isspace() and not is_arabic_char(t[0]) and not t[0].isdigit()
                      and not all(is_mark(ch) for ch in t))
        if not suspicious:
            k += 1
            continue
        j = k
        while j + 1 < len(seq) and seq[j + 1][0] - seq[j][1] < 0.05 * size and \
                not seq[j + 1][2].isspace() and not is_arabic_char(seq[j + 1][2][0]) and \
                not seq[j + 1][2][0].isdigit():
            j += 1
        if k > 0 and j + 1 < len(seq):
            L, R = seq[k - 1], seq[j + 1]
            tight = (x0 - L[1]) < 0.05 * size and (R[0] - seq[j][1]) < 0.05 * size
            if tight and arabic_letter(L[2]) and arabic_letter(R[2]) and (j - k) < 4:
                for m in range(k, j + 1):
                    drop.add(m)
                bridge.append((k - 1, seq[j][1]))
        k = j + 1
    for li, newx1 in bridge:
        it = seq[li]
        seq[li] = (it[0], newx1, it[2], it[3], it[4])
    return [it for m, it in enumerate(seq) if m not in drop]


_TANWEEN = "\u064b\u064c\u064d"


def _clean_arabic(text: str) -> str:
    """Remove kashida and impossible tanween (tanween is always word-final)."""
    text = text.replace("\u0640", "")
    out = []
    n = len(text)
    for i, ch in enumerate(text):
        if ch in _TANWEEN:
            # next base letter after this mark cluster
            j = i + 1
            while j < n and is_mark(text[j]):
                j += 1
            nxt = text[j] if j < n else " "
            if is_arabic_char(nxt) and nxt.isalpha():
                # allowed: fathatan before a word-final alef / alef maqsura
                k = j + 1
                while k < n and is_mark(text[k]):
                    k += 1
                after = text[k] if k < n else " "
                if not (ch == "\u064b" and nxt in "\u0627\u0649" and not (is_arabic_char(after) and after.isalpha())):
                    continue
        out.append(ch)
    return "".join(out)


def _walk_chars(container, out):
    from pdfminer.layout import LTChar

    for obj in container:
        if isinstance(obj, LTChar):
            obj.waraq_seq = len(out)  # drawing order (pdfminer keeps it inside a line)
            out.append(obj)
        elif hasattr(obj, "__iter__") and not isinstance(obj, str):
            try:
                _walk_chars(obj, out)
            except TypeError:
                pass
    return out


def _has_image(container) -> bool:
    from pdfminer.layout import LTImage, LTFigure

    for obj in container:
        if isinstance(obj, LTImage):
            return True
        if isinstance(obj, LTFigure) and _has_image(obj):
            return True
    return False


class _RawLine:
    __slots__ = ("chars", "y0", "y1")

    def __init__(self, c):
        self.chars = [c]
        self.y0, self.y1 = c.y0, c.y1

    def accepts(self, c) -> bool:
        ov = min(self.y1, c.y1) - max(self.y0, c.y0)
        return ov > 0.55 * min(self.y1 - self.y0, c.y1 - c.y0)


def _cluster_lines(chars) -> List[_RawLine]:
    """Group glyphs into visual lines by vertical overlap (producer independent)."""
    chars = [c for c in chars if c.y1 - c.y0 > 0.5 and not c.get_text().startswith("(cid:")
             or (c.get_text().startswith("(cid:") and False)]
    chars.sort(key=lambda c: -(c.y0 + c.y1) / 2)
    lines: List[_RawLine] = []
    for c in chars:
        # only compare with recent lines (sorted top->bottom)
        for ln in reversed(lines[-6:]):
            if ln.accepts(c):
                ln.chars.append(c)
                break
        else:
            lines.append(_RawLine(c))
    lines.sort(key=lambda l: -(l.y0 + l.y1) / 2)
    return lines


def _segments(chars, gap_factor: float = 2.5):
    """Split one visual line at big horizontal gaps (table cells / columns)."""
    cs = sorted(chars, key=lambda c: c.x0)
    if not cs:
        return []
    size = statistics.median([c.size for c in cs]) or 10
    segs = [[cs[0]]]
    right = cs[0].x1
    for c in cs[1:]:
        if c.x0 - right > gap_factor * size:
            segs.append([c])
        else:
            segs[-1].append(c)
        right = max(right, c.x1)
    return segs


def extract_pages(path, password: Optional[str] = None, pages: Optional[Iterable[int]] = None,
                  progress: Optional[Callable[[int, int], None]] = None) -> List[PageText]:
    """Return logical text per page (``pages`` are 0-based indexes)."""
    from pdfminer.high_level import extract_pages as pm_pages
    from pdfminer.layout import LAParams
    from pdfminer.pdfdocument import PDFPasswordIncorrect
    from pdfminer.pdfparser import PDFSyntaxError

    from .util import open_reader

    n = len(open_reader(path, password).pages)  # validates the password early
    wanted = sorted(pages) if pages is not None else list(range(n))
    laparams = LAParams(line_margin=0.4, char_margin=3.0, word_margin=0.1)
    out: List[PageText] = []
    try:
        it = pm_pages(str(path), password=password or "", page_numbers=wanted, laparams=laparams)
        for k, (layout, pno) in enumerate(zip(it, wanted), 1):
            out.append(_page_from_layout(layout, pno))
            if progress:
                progress(k, len(wanted))
    except PDFPasswordIncorrect:
        raise WrongPasswordError(str(path)) if password else NeedsPasswordError(str(path))
    except PDFSyntaxError as e:
        raise WaraqError(f"Not a valid PDF: {e}", "bad_pdf")
    return out


def _page_from_layout(layout, pno: int) -> PageText:
    pt = PageText(number=pno + 1, has_images=_has_image(layout))
    chars = _walk_chars(layout, [])
    letters = [c.get_text() for c in chars]
    ar = sum(1 for t in letters for ch in t if is_arabic_char(ch))
    la = sum(1 for t in letters for ch in t if ch.isascii() and ch.isalpha())
    page_rtl = ar > 0 and ar >= 0.3 * (ar + la)
    raw_lines: List[Line] = []
    for rl in _cluster_lines(chars):
        segs = _segments(rl.chars)
        if page_rtl:
            segs = list(reversed(segs))
        for seg in segs:
            units, size = _line_units(seg)
            if not units:
                continue
            raw = "".join(units)
            rtl = has_arabic(raw)
            text = _clean_arabic(visual_units_to_logical(units, base="R" if page_rtl else None)) if rtl else raw
            pt.chars += sum(len(u) for u in units if not u.isspace())
            raw_lines.append(Line(text, min(c.x0 for c in seg), rl.y0, max(c.x1 for c in seg), rl.y1,
                                  size, rtl))
    # group lines into blocks (paragraphs): a new block starts after a bigger
    # vertical gap than usual, a font size change, or a line that ends early
    # (the last line of a justified/aligned paragraph)
    pitches = []
    for a, b in zip(raw_lines, raw_lines[1:]):
        d = ((a.y0 + a.y1) - (b.y0 + b.y1)) / 2
        if d > 0.3 * max(a.size, 1):
            pitches.append(d)
    usual = statistics.median(pitches) if pitches else 0
    left = min((l.x0 for l in raw_lines), default=0)
    right = max((l.x1 for l in raw_lines), default=0)
    width = max(right - left, 1)
    blocks: List[Block] = []
    prev = None
    for ln in raw_lines:
        new_block = prev is None
        if prev is not None:
            pitch = ((prev.y0 + prev.y1) - (ln.y0 + ln.y1)) / 2
            if pitch < 0.3 * max(ln.size, 1):
                new_block = True                      # same row (table cell / column)
            elif usual and pitch > 1.3 * usual:
                new_block = True
            elif abs(ln.size - prev.size) > 1.5:
                new_block = True
            else:
                if prev.rtl or page_rtl:
                    ended = (prev.x0 - left) > 0.12 * width
                else:
                    ended = (right - prev.x1) > 0.12 * width
                new_block = ended
        if new_block:
            blocks.append(Block([ln], ln.rtl or page_rtl))
        else:
            blocks[-1].lines.append(ln)
        prev = ln
    for b in blocks:
        b.rtl = any(l.rtl for l in b.lines) and page_rtl or all(l.rtl for l in b.lines)
    pt.blocks = blocks
    return pt


def pages_to_text(pages: List[PageText], page_markers: bool = False) -> str:
    parts = []
    for p in pages:
        if page_markers:
            parts.append(f"--- {p.number} ---")
        parts.extend(p.lines())
        parts.append("")
    return "\n".join(parts).rstrip() + "\n"


def save_txt(pages: List[PageText], dst, page_markers: bool = False) -> None:
    from .util import safe_write_bytes

    # UTF-8 with BOM so old Notepad versions detect Arabic correctly
    safe_write_bytes(dst, pages_to_text(pages, page_markers).encode("utf-8-sig"))


def save_docx(pages: List[PageText], dst, layout: str = "paragraphs", font: str = "Arial",
              size_pt: float = 14) -> None:
    """Write a right-to-left Word document.

    layout="paragraphs": one Word paragraph per text block (lines joined,
    easy to edit).  layout="lines": one paragraph per PDF line.
    """
    import io

    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    from .util import safe_write_bytes

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size_pt)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), font)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(size_pt * 2)))
    rpr.append(szcs)

    def add_par(text: str, rtl: bool):
        p = doc.add_paragraph()
        if rtl:
            ppr = p._p.get_or_add_pPr()
            bidi = OxmlElement("w:bidi")
            ppr.insert(0, bidi)
        run = p.add_run(text)
        if rtl:
            r = run._r.get_or_add_rPr()
            r.append(OxmlElement("w:rtl"))
        return p

    first = True
    for pg in pages:
        if not first:
            doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE) if doc.paragraphs else None
        first = False
        for b in pg.blocks:
            if layout == "lines":
                for l in b.lines:
                    add_par(l.text, l.rtl or b.rtl)
            else:
                add_par(b.text, b.rtl)
        if not pg.blocks:
            add_par("", True)
    buf = io.BytesIO()
    doc.save(buf)
    safe_write_bytes(dst, buf.getvalue())


def docx_paragraph_texts(path) -> List[str]:
    from docx import Document

    return [p.text for p in Document(str(path)).paragraphs]
