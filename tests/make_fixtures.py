"""Generate realistic test PDFs into tests/_fixtures (not committed).

Three different "producers" are used for Arabic, because each writes
Arabic differently inside the PDF:
  * reportlab  - pre-shaped presentation forms in visual order (typical of
                 many generators and old tools)
  * chrome     - Skia/HarfBuzz shaping (like browsers' "Save as PDF")
  * writer     - LibreOffice Writer export (closest to Word "Save as PDF")
"""
from __future__ import annotations

import io
import os
import random
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUT = ROOT / "tests" / "_fixtures"
DATA = ROOT / "tests" / "data"
AMIRI = ROOT / "waraq" / "assets" / "fonts" / "Amiri-Regular.ttf"


def arabic_lines():
    return [l for l in (DATA / "arabic_lines.txt").read_text(encoding="utf-8").splitlines() if l.strip()]


def english_lines():
    return [l for l in (DATA / "english_lines.txt").read_text(encoding="utf-8").splitlines() if l.strip()]


def make_reportlab_arabic(path: Path, lines, pages=2):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    from waraq.core.arabic import shape_for_display

    pdfmetrics.registerFont(TTFont("Amiri", str(AMIRI)))
    c = canvas.Canvas(str(path), pagesize=(595, 842))
    for p in range(pages):
        y = 780
        c.setFont("Amiri", 16)
        for ln in lines:
            c.drawRightString(555, y, shape_for_display(ln))
            y -= 34
        c.setFont("Helvetica", 9)
        c.drawString(40, 30, f"page {p + 1}")
        c.showPage()
    c.save()


HTML = """<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8">
<style>@page {{ size: A4; margin: 2cm; }}
body {{ font-family: '{font}'; font-size: 17px; line-height: 2.1; }}
p {{ margin: 0; white-space: nowrap; }}</style></head><body>{body}</body></html>"""


def make_chrome_arabic(path: Path, lines, font="Noto Naskh Arabic"):
    chrome = shutil.which("google-chrome") or shutil.which("chromium")
    if not chrome:
        return False
    body = "".join(f"<p>{l}</p>" for l in lines)
    with tempfile.TemporaryDirectory() as td:
        html = Path(td) / "a.html"
        html.write_text(HTML.format(font=font, body=body), encoding="utf-8")
        cmd = [chrome, "--headless=new", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
               f"--user-data-dir={td}/prof", f"--print-to-pdf={path}", html.as_uri()]
        try:
            subprocess.run(cmd, check=True, capture_output=True, timeout=40)
        except (subprocess.TimeoutExpired, subprocess.CalledProcessError):
            return False
    return path.exists()


def make_writer_arabic(path: Path, lines, font="Amiri", justify=False, size=15):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    for l in lines:
        p = doc.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        ppr.insert(0, OxmlElement("w:bidi"))
        if justify:
            jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "both"); ppr.append(jc)
        r = p.add_run(l)
        rpr = r._r.get_or_add_rPr()
        fonts = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(a), font)
        rpr.append(fonts)
        rpr.append(OxmlElement("w:rtl"))
        r.font.size = Pt(size)
        sz = OxmlElement("w:szCs"); sz.set(qn("w:val"), str(size * 2)); rpr.append(sz)
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / (path.stem + ".docx")
        doc.save(src)
        (Path(td) / ".fonts").mkdir()
        shutil.copy(AMIRI, Path(td) / ".fonts" / AMIRI.name)
        env = dict(os.environ, HOME=td)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", td, str(src)],
                       check=True, capture_output=True, timeout=180, env=env)
        shutil.copy(Path(td) / (path.stem + ".pdf"), path)
    return True


def make_english(path: Path, pages=50):
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=(612, 792))
    for p in range(pages):
        c.setFont("Helvetica-Bold", 28)
        c.drawString(72, 700, f"Page {p + 1}")
        c.setFont("Helvetica", 12)
        y = 660
        for ln in english_lines():
            c.drawString(72, y, ln)
            y -= 18
        c.showPage()
    c.save()


def photo(w, h, seed):
    """A photo-like image (gradients, shapes, noise) that JPEG cannot squash trivially."""
    from PIL import Image, ImageDraw, ImageFilter

    rnd = random.Random(seed)
    im = Image.new("RGB", (w, h))
    px = Image.linear_gradient("L").resize((w, h))
    r = px
    g = px.rotate(90).resize((w, h))
    b = Image.radial_gradient("L").resize((w, h))
    im = Image.merge("RGB", (r, g, b))
    d = ImageDraw.Draw(im)
    for _ in range(60):
        x0, y0 = rnd.randrange(w), rnd.randrange(h)
        rad = rnd.randrange(20, max(21, w // 6))
        d.ellipse([x0, y0, x0 + rad, y0 + rad], fill=tuple(rnd.randrange(256) for _ in range(3)))
    im = im.filter(ImageFilter.GaussianBlur(3))
    noise = Image.effect_noise((w, h), 40).convert("RGB")
    return Image.blend(im, noise, 0.18)


def make_photos_pdf(path: Path, pages=6):
    """Large 'scanned brochure': big lossless (Flate) photos + a little text."""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=(595, 842))
    for p in range(pages):
        im = photo(2480, 1754, p)
        c.drawImage(ImageReader(im), 20, 300, 555, 392)
        c.setFont("Helvetica", 14)
        c.drawString(40, 200, f"Photo page {p + 1} - text must survive compression")
        c.showPage()
    c.save()


def make_scanned_arabic(src_pdf: Path, path: Path, dpi=200):
    """Rasterise a text PDF to images (like a scanner) and wrap them in a PDF."""
    import pypdfium2 as pdfium
    from PIL import ImageFilter
    from waraq.core.images import images_to_pdf

    doc = pdfium.PdfDocument(str(src_pdf))
    tmp = Path(tempfile.mkdtemp())
    imgs = []
    for i in range(len(doc)):
        im = doc[i].render(scale=dpi / 72).to_pil().convert("L")
        im = im.rotate(0.4, expand=False, fillcolor=255).filter(ImageFilter.GaussianBlur(0.6))
        p = tmp / f"scan_{i}.png"
        im.save(p, dpi=(dpi, dpi))
        imgs.append(p)
    images_to_pdf(imgs, path, page_size="fit")
    shutil.rmtree(tmp, ignore_errors=True)


def make_images(dirp: Path):
    from PIL import Image

    dirp.mkdir(parents=True, exist_ok=True)
    photo(1200, 800, 1).save(dirp / "landscape.jpg", quality=90)
    photo(600, 900, 2).save(dirp / "portrait.png")
    # a JPEG whose EXIF says "rotate 90 degrees" (phone photo)
    im = photo(800, 400, 3)
    exif = Image.Exif()
    exif[0x0112] = 6
    im.save(dirp / "phone_rotated.jpg", quality=90, exif=exif)
    rgba = Image.new("RGBA", (500, 300), (255, 0, 0, 0))
    rgba.paste((0, 128, 0, 255), (100, 100, 400, 200))
    rgba.save(dirp / "transparent.png")


def make_signature(path: Path):
    from PIL import Image, ImageDraw

    im = Image.new("RGB", (600, 220), (250, 250, 248))
    d = ImageDraw.Draw(im)
    pts = [(40, 150), (90, 60), (140, 170), (200, 50), (260, 160), (330, 70), (400, 150), (560, 100)]
    d.line(pts, fill=(10, 20, 120), width=9, joint="curve")
    d.line([(60, 190), (540, 185)], fill=(10, 20, 120), width=5)
    im.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    lines = arabic_lines()
    make_reportlab_arabic(OUT / "arabic_reportlab.pdf", lines)
    ok_c = make_chrome_arabic(OUT / "arabic_chrome.pdf", lines) if os.environ.get("WARAQ_TRY_CHROME") else False
    ok_w = make_writer_arabic(OUT / "arabic_writer.pdf", lines, "Amiri")
    make_writer_arabic(OUT / "arabic_writer_dejavu.pdf", lines, "DejaVu Sans")
    paras = [l for l in (DATA / "arabic_paragraphs.txt").read_text(encoding="utf-8").splitlines() if l.strip()]
    make_writer_arabic(OUT / "arabic_paragraphs.pdf", paras, "Amiri", justify=True, size=14)
    make_english(OUT / "english_50p.pdf")
    make_photos_pdf(OUT / "photos_large.pdf")
    make_scanned_arabic(OUT / "arabic_reportlab.pdf", OUT / "arabic_scanned.pdf")
    make_images(OUT / "images")
    make_signature(OUT / "signature.jpg")
    from waraq.core.security import protect
    protect(OUT / "english_50p.pdf", OUT / "encrypted_user123.pdf", "user123", "owner456")
    # a page with /Rotate 90 for stamping tests
    from waraq.core.pages import rotate
    rotate(OUT / "english_50p.pdf", OUT / "rotated90.pdf", 90, pages="1-3")
    (OUT / "not_a_pdf.pdf").write_bytes(b"hello, I am not a PDF")
    print("chrome:", ok_c, "writer:", ok_w)
    for p in sorted(OUT.rglob("*")):
        if p.is_file():
            print(f"{p.relative_to(OUT)}  {p.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
